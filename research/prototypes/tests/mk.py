import sys, os
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
out=sys.argv[1]
doc=Document()
for t in ["السؤال ١: ما الهدف الرئيس من التوجيه الطلابي في المدرسة؟",
          "أ) تنمية شخصية الطالب في الجوانب النفسية والاجتماعية عام 2024",
          "ب) إعداد الجداول الدراسية (Schedules) للمعلمين"]:
    p=doc.add_paragraph()
    pPr=p._p.get_or_add_pPr(); bidi=pPr.makeelement(qn('w:bidi'),{}); pPr.append(bidi)
    r=p.add_run(t); r.font.size=Pt(14); r.font.name='DejaVu Sans'
    rPr=r._r.get_or_add_rPr(); rtl=rPr.makeelement(qn('w:rtl'),{}); rPr.append(rtl)
doc.save(os.path.join(out,'q.docx'))
# reportlab PDF using reshaper+bidi (visual order, presentation forms) -- common in generated PDFs
from reportlab.pdfgen import canvas
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics
import arabic_reshaper
from bidi.algorithm import get_display
pdfmetrics.registerFont(TTFont('DV','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
c=canvas.Canvas(os.path.join(out,'visual.pdf'))
c.setFont('DV',14)
c.drawRightString(550,800,get_display(arabic_reshaper.reshape("ما الهدف الرئيس من التوجيه الطلابي في المدرسة؟")))
c.save()
